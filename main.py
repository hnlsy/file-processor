#!/usr/bin/env python3
"""File Processor - 批量文件处理工具"""

import os
import glob
import argparse
from pathlib import Path

import openpyxl
from pypdf import PdfReader, PdfWriter
from PIL import Image


class ExcelProcessor:
    """Excel 批量处理"""

    @staticmethod
    def merge_sheets(file_path: str) -> list:
        """合并 Excel 所有 Sheet"""
        wb = openpyxl.load_workbook(file_path, read_only=True)
        all_rows = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                all_rows.append(row)
        wb.close()
        return all_rows

    @staticmethod
    def batch_read(folder: str) -> dict:
        """批量读取文件夹内所有 Excel"""
        results = {}
        for f in glob.glob(os.path.join(folder, "*.xlsx")):
            wb = openpyxl.load_workbook(f, read_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    rows.append(row)
            wb.close()
            results[os.path.basename(f)] = rows
            print(f"  已读取: {os.path.basename(f)} ({len(rows)} 行)")
        return results

    @staticmethod
    def clean_and_export(input_file: str, output_file: str):
        """清洗数据并导出"""
        wb = openpyxl.load_workbook(input_file)
        ws = wb.active
        cleaned = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    cell.value = cell.value.strip()
                    cleaned += 1
        wb.save(output_file)
        print(f"  清洗完成: {cleaned} 个单元格 -> {output_file}")


class PdfProcessor:
    """PDF 批量处理"""

    @staticmethod
    def extract_text(file_path: str) -> str:
        """提取 PDF 文本"""
        reader = PdfReader(file_path)
        text = ""
        for i, page in enumerate(reader.pages):
            text += f"--- 第 {i+1} 页 ---\n"
            text += page.extract_text() + "\n"
        return text

    @staticmethod
    def merge_pdfs(pdf_list: list, output_path: str):
        """合并多个 PDF"""
        writer = PdfWriter()
        for pdf_path in pdf_list:
            reader = PdfReader(pdf_path)
            for page in reader.pages:
                writer.add_page(page)
            print(f"  已合并: {os.path.basename(pdf_path)}")
        with open(output_path, "wb") as f:
            writer.write(f)
        print(f"  输出: {output_path}")


class ImageProcessor:
    """图片批量处理"""

    @staticmethod
    def batch_compress(folder: str, quality: int = 85):
        """批量压缩图片"""
        output_dir = os.path.join(folder, "compressed")
        os.makedirs(output_dir, exist_ok=True)
        for f in glob.glob(os.path.join(folder, "*.{jpg,jpeg,png}", )) + glob.glob(os.path.join(folder, "*.jpg")) + glob.glob(os.path.join(folder, "*.png")):
            try:
                img = Image.open(f)
                out_path = os.path.join(output_dir, os.path.basename(f))
                img.save(out_path, quality=quality, optimize=True)
                orig = os.path.getsize(f) / 1024
                comp = os.path.getsize(out_path) / 1024
                print(f"  {os.path.basename(f)}: {orig:.0f}KB -> {comp:.0f}KB")
            except Exception as e:
                print(f"  跳过 {os.path.basename(f)}: {e}")

    @staticmethod
    def batch_resize(folder: str, width: int, height: int):
        """批量调整尺寸"""
        output_dir = os.path.join(folder, "resized")
        os.makedirs(output_dir, exist_ok=True)
        for f in glob.glob(os.path.join(folder, "*.jpg")) + glob.glob(os.path.join(folder, "*.png")):
            try:
                img = Image.open(f)
                img = img.resize((width, height), Image.Resampling.LANCZOS)
                out_path = os.path.join(output_dir, os.path.basename(f))
                img.save(out_path)
                print(f"  {os.path.basename(f)}: -> {width}x{height}")
            except Exception as e:
                print(f"  跳过 {os.path.basename(f)}: {e}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="File Processor - 批量文件处理工具")
    sub = parser.add_subparsers(dest="command")

    # Excel 命令
    excel_cmd = sub.add_parser("excel", help="Excel 处理")
    excel_cmd.add_argument("--merge", help="合并 Excel 所有 Sheet")
    excel_cmd.add_argument("--batch", help="批量读取文件夹")
    excel_cmd.add_argument("--clean", nargs=2, help="清洗并导出: input output")

    # PDF 命令
    pdf_cmd = sub.add_parser("pdf", help="PDF 处理")
    pdf_cmd.add_argument("--extract", help="提取 PDF 文本")
    pdf_cmd.add_argument("--merge", nargs="+", help="合并多个 PDF")
    pdf_cmd.add_argument("--output", help="输出文件名")

    # Image 命令
    img_cmd = sub.add_parser("image", help="图片处理")
    img_cmd.add_argument("--compress", help="批量压缩图片文件夹")
    img_cmd.add_argument("--resize", nargs=3, help="调整尺寸: folder width height")

    args = parser.parse_args()

    if args.command == "excel":
        if args.merge:
            rows = ExcelProcessor.merge_sheets(args.merge)
            print(f"合并完成: {len(rows)} 行")
        elif args.batch:
            ExcelProcessor.batch_read(args.batch)
        elif args.clean:
            ExcelProcessor.clean_and_export(args.clean[0], args.clean[1])
    elif args.command == "pdf":
        if args.extract:
            text = PdfProcessor.extract_text(args.extract)
            print(text[:500])
        elif args.merge:
            output = args.output or "merged.pdf"
            PdfProcessor.merge_pdfs(args.merge, output)
    elif args.command == "image":
        if args.compress:
            ImageProcessor.batch_compress(args.compress)
        elif args.resize:
            ImageProcessor.batch_resize(args.resize[0], int(args.resize[1]), int(args.resize[2]))
    else:
        parser.print_help()
